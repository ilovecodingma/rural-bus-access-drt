# -*- coding: utf-8 -*-
"""분석보고서(docx) 생성. 사용: python 03_build_report.py <out폴더> <결과 docx 경로> [코드 공개 URL]"""
import sys, json
sys.stdout.reconfigure(encoding='utf-8')
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]; DST = sys.argv[2]
URL = sys.argv[3] if len(sys.argv) > 3 else '[GitHub 업로드 후 URL 기재]'
R = json.load(open(f'{OUT}/results.json', encoding='utf-8'))
T = pd.read_csv(f'{OUT}/table_type.csv').set_index('유형')
Bd = pd.read_csv(f'{OUT}/table_band.csv')
G = pd.read_csv(f'{OUT}/table_sigungu.csv')
S = pd.read_csv(f'{OUT}/table_sido.csv')
V = R['vul']; sens = R['sens']
f0 = lambda x: f'{x:,.0f}'
pc = lambda x, d=1: f'{x * 100:.{d}f}%'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.9); sec.top_margin = Cm(1.7); sec.bottom_margin = Cm(1.6)
FONT = 'Malgun Gothic'
st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(9.3)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
st.paragraph_format.space_after = Pt(1.5); st.paragraph_format.line_spacing = 1.12
NAVY = RGBColor(0x1F, 0x3A, 0x5F)

def para(text='', bold=False, size=None, color=None, align=None, after=None, before=None, indent=None, keep=False):
    p = doc.add_paragraph()
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    if indent is not None: p.paragraph_format.left_indent = Cm(indent); p.paragraph_format.first_line_indent = Cm(-0.35)
    if keep: p.paragraph_format.keep_with_next = True
    # **굵게** 구간 지원
    parts = text.split('**')
    for i, t in enumerate(parts):
        if not t: continue
        r = p.add_run(t); r.bold = bold or (i % 2 == 1)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        r.font.name = FONT; r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    return p
def h1(t):
    p = para(t, bold=True, size=11.5, color=NAVY, before=7, after=2, keep=True)
    pb = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '6'), ('space', '1'), ('color', '1F3A5F')): b.set(qn('w:' + k), v)
    pb.append(b); p._p.get_or_add_pPr().append(pb)
def h2(t): para(t, bold=True, size=9.8, before=3, after=1, keep=True)
def b1(t): para('○ ' + t, indent=0.45, after=1)
def b2(t): para('- ' + t, indent=0.95, after=0.5)
def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcol); tcPr.append(sh)
def table(rows, widths, size=8.2, header=True, align_num=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ''
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(v)); r.font.size = Pt(size); r.font.name = FONT; r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
            if header and i == 0: r.bold = True; shade(c, 'DCE6F1'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif align_num and j > 0 and str(v)[:1].isdigit(): p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return t
def fig(path, w, cap, after=3):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(w))
    para(cap, size=8, color=RGBColor(0x55, 0x55, 0x55), align=WD_ALIGN_PARAGRAPH.CENTER, after=after)

# ================= 표지 =================
for _ in range(7): para('')
para('AI와 함께하는 교통문제 해결을 위한 데이터 분석 공모전', size=12, color=RGBColor(0x55, 0x55, 0x55), align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
para('분 석 보 고 서', bold=True, size=26, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=22)
para('고령 농촌 읍면동의 버스정류장 접근성 격차 분석과\n수요응답형 교통(DRT) 우선 도입 지역 제안', bold=True, size=15,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
para('- 전국 3,042개 읍면동 단위 공공데이터 결합 분석 -', size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=60)
table([['이름/팀명', '박성상 (개인)'], ['분석 기준', '정류장 2025.10.31 / 주민등록 인구 2025.10 / 행정동 경계 2025.10.01'],
       ['분석 코드', URL]], [3.2, 12.4], size=9.5, header=False, align_num=False)
para('')
para('2026. 10.', size=11, align=WD_ALIGN_PARAGRAPH.CENTER, before=40)
doc.add_page_break()

# ================= 본문 =================
h1('1. 주제')
para(f'**고령 농촌 읍면동의 버스정류장 접근성 격차 분석 - 전국 읍면동 {R["n"]:,}곳의 정류장 밀도와 고령비중을 결합한 이동취약 지역 식별 및 DRT 우선 도입 제안**', after=2)

h1('2. 분석 배경 및 목적')
b1('교통은 일자리·교육·의료 등 삶의 기회로 연결되는 통로이며, 공고문이 지적하듯 대중교통이 부족한 지역의 이동 격차는 여전히 존재함. 정책 방향도 자동차 소통 중심에서 사람 중심·사회적 형평성 중심으로 이동 중.')
b1('수송부문 온실가스의 약 97%가 도로교통에서 발생(공고문 인용). 대중교통 접근이 어려운 지역은 승용차 이용 외 대안이 없어 탄소중립과 이동권이 함께 걸린 문제임.')
b1(f'기존 평가는 시군구 평균 또는 인구 대비 정류장 수 위주여서, 면적이 넓고 거주지가 흩어진 농촌의 격차가 가려질 수 있음. 본 분석에서 **면 지역은 고령자 1천 명당 정류장 수가 동 지역의 {T.loc["면","고령자천명당정류장"] / T.loc["동","고령자천명당정류장"]:.1f}배**({T.loc["면","고령자천명당정류장"]:.1f}개 대 {T.loc["동","고령자천명당정류장"]:.1f}개)로 오히려 많게 나타나지만, **면적 기준으로는 정류장 1개가 담당하는 면적이 {T.loc["면","정류장당면적ha"] / T.loc["동","정류장당면적ha"]:.1f}배**({T.loc["면","정류장당면적ha"]:.0f}ha 대 {T.loc["동","정류장당면적ha"]:.0f}ha) 넓음. 어느 지표를 쓰느냐에 따라 결론이 뒤집힘.')
b1('목적: (1) 읍면동 단위로 고령비중과 정류장 밀도를 결합해 "이동취약 읍면동"을 식별, (2) 규모·분포·집중 지역을 계량, (3) 정류장 신설이 아닌 수요응답형 교통(DRT) 중심의 우선 도입 지역과 운영 단위를 제안.')

h1('3. 활용 데이터')
table([['활용 데이터(명)', '제공기관(명)', '출처 플랫폼(명)', 'URL / 기준 / 선정 이유'],
       ['전국 버스정류장 위치정보 (정류장 227,053개, 위·경도)', '국토교통부', '공공데이터포털 (파일데이터, CSV)',
        'https://www.data.go.kr/data/15067528/fileData.do · 정보수집일 2025.10.31 · 이용허락범위 제한 없음 · 전국 단일 기준으로 정류장 위치 비교 가능'],
       ['연령별 인구현황 (읍면동, 1세 단위, 2025년 10월)', '행정안전부', '주민등록 인구통계 (jumin.mois.go.kr)',
        'https://jumin.mois.go.kr/ageStatMonth.do · 2025.10 · 읍면동 단위 65세 이상 인구 산출'],
       ['행정동 경계 GeoJSON (2025.10.01판)', '통계청 SGIS 원자료를 가공한 vuski/admdongkor', 'GitHub (CC BY 4.0)',
        'https://github.com/vuski/admdongkor/tree/master/ver20251001 · 정류장 공간결합·면적 산출용']],
      [4.4, 2.7, 3.2, 7.3], size=7.8, align_num=False)
para('※ 모두 로그인·API 키 없이 공개된 자료이며 출처·기준시점·이용조건을 위와 같이 명시함. 개인정보·비공개 자료는 사용하지 않음.', size=8, after=1, before=1)

h1('4. 분석 과정 및 방법')
b1('**수집**: 위 3종을 스크립트(00_download.py)로 직접 내려받음. 정류장·인구 파일은 재다운로드 결과가 최초 파일과 바이트 단위로 동일함을 확인.')
b1(f'**공간결합**: 정류장 227,053개 중 226,882개(99.92%)를 행정동 경계에 포인트-인-폴리곤으로 배정하고, 읍면동 이름을 정규화(공백·가운뎃점 제거)해 인구표와 결합. 결합 후 읍면동 단위 {R["n_all"]:,}개 확보(이름 표기 불일치로 인구 단위 36개·정류장 3,409개는 결합에서 제외).')
b1('**지표**: 65세 이상 비중 = 65세 이상 인구/전체 인구, 정류장 밀도 = 정류장 수/면적(㎢, 경계 면적을 위도 보정 근사). 읍면동 유형(동·읍·면)은 명칭으로 구분.')
b1(f'**수집 누락 의심 지역 제외**: 정류장 수가 비정상적으로 적은 시군구(예: 고성군 0개, 양구군 2개)는 자료 누락으로 판단해 분석에서 제외({len(R["excluded_sgg"])}개 시군구: 강릉·동해·삼척·속초·고성·양구·인제·정선·평창·화천·강진·보성·영광·익산·증평). 판정 규칙은 코드에 공개(02_analysis.py). 제외 후 분석 대상은 읍면동 **{R["n"]:,}개**, 인구 {f0(R["pop"])}명, 정류장 {f0(R["stops"])}개.')
b1('**이동취약 읍면동 정의**: 65세 이상 비중 40% 이상(초고령사회 기준 20%의 2배) **이면서** 정류장 밀도가 반경 500m 가정 서비스권 1개의 면적(π×0.5²=0.785㎢)당 1개 미만, 즉 1.27개/㎢ 미만. 반경은 분석자 가정이므로 300m·400m로 민감도 분석.')
b1('**기법**: 기술통계, 유형별 비교, 고령비중 구간별 비교, Spearman 순위상관, 시도·시군 집계, 임계값 민감도 분석(pandas, shapely, scipy, matplotlib).')

h1('5. AI 기술 활용')
b1('**사용 서비스**: Anthropic Claude (Claude Code, 모델 Claude Opus 5.5).')
b1('**활용 범위**: 공고·양식 확인, 공개 데이터 탐색과 내려받기 스크립트 작성, 전처리·공간결합·통계 분석 코드 작성 및 실행, 시각화, 보고서 초안 작성. 본문 수치는 모두 위 코드가 공개 데이터에서 계산한 값이며 AI가 생성한 데이터는 없음.')
b1('**주요 프롬프트(요지)**: (1) "교통문제 해결 데이터 분석 공모전 공고와 양식을 확인하고, 공개 데이터를 직접 수집해 분석 보고서를 작성하라." (2) "전국 정류장 위치와 읍면동 연령별 인구를 결합해 고령 인구 대비 대중교통 접근이 취약한 지역을 찾아라. 수치는 실제 데이터에서 계산한 값만 써라." (3) "공고 양식 항목(AI 활용 포함)에 맞추어 5쪽 내외 보고서로 정리하라."')
b1('**검증**: 정류장 공간배정률(99.92%), 인구 총계 대조, 임계값 민감도, 자료 누락 의심 지역 제외 규칙, 재다운로드 동일성을 확인. 입력 데이터는 모두 공개 자료로, 개인정보·비공개 자료는 AI 서비스에 입력하지 않음.')

h1('6. 분석 내용 및 결과')
h2('가. 읍·면·동 유형별 격차')
fig(f'{OUT}/fig2_type.png', 13.5, f'[그림 1] 유형별 정류장 밀도·고령비중·정류장당 면적 (읍면동 {R["n"]:,}개)', after=1)
rows = [['유형', '읍면동 수', '인구(명)', '65세 이상 비중', '정류장 수', '밀도(개/㎢)', '정류장당 면적(ha)', '고령자 1천 명당 정류장', '밀도 1.27 미만 비율']]
for k in ['동', '읍', '면']:
    r = T.loc[k]; rows.append([k, f0(r['읍면동수']), f0(r['인구']), pc(r['고령비중']), f0(r['정류장']), f'{r["밀도"]:.2f}',
                               f'{r["정류장당면적ha"]:.1f}', f'{r["고령자천명당정류장"]:.1f}', pc(r['밀도기준미만비율'])])
table(rows, [1.2, 1.6, 2.0, 1.9, 1.7, 1.7, 2.2, 2.6, 2.2], size=7.8)
para(f'- 면 지역은 인구의 {pc(T.loc["면","인구"] / T["인구"].sum())}를 차지하나 65세 이상 비중이 {pc(T.loc["면","고령비중"])}(동 {pc(T.loc["동","고령비중"])})이고, 정류장 밀도는 동의 {T.loc["면","밀도"] / T.loc["동","밀도"] * 100:.0f}% 수준({T.loc["면","밀도"]:.2f}개/㎢). 면 지역 중 {pc(T.loc["면","밀도기준미만비율"], 0)}가 밀도 기준 미만인 반면 동 지역은 {pc(T.loc["동","밀도기준미만비율"])}에 그침.', size=8.8, before=2)
b = Bd.set_index('band')
para(f'- 고령비중 구간별로 보면 20% 미만 읍면동의 정류장 밀도는 {b.loc["20% 미만","dens"]:.2f}개/㎢, 50% 이상은 {b.loc["50% 이상","dens"]:.2f}개/㎢로 **{b.loc["20% 미만","dens"] / b.loc["50% 이상","dens"]:.1f}배 차이**. 50% 이상 구간의 {pc(b.loc["50% 이상","low"], 0)}가 밀도 기준 미만. 읍면동별 고령비중과 정류장 밀도의 Spearman 상관계수는 **{R["rho"]:.2f}**(강한 음의 상관).', size=8.8)

h2('나. 이동취약 읍면동의 규모와 분포')
fig(f'{OUT}/fig1_scatter.png', 11.6, '[그림 2] 고령비중과 정류장 밀도 (주황: 이동취약 읍면동)', after=1)
# 지도 + 요약표 나란히
t = doc.add_table(rows=1, cols=2); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
lc, rc = t.cell(0, 0), t.cell(0, 1); lc.width = Cm(8.2); rc.width = Cm(9.0)
lp = lc.paragraphs[0]; lp.alignment = WD_ALIGN_PARAGRAPH.CENTER; lp.add_run().add_picture(f'{OUT}/fig3_map.png', width=Cm(7.0))
cp = lc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
cr = cp.add_run('[그림 3] 이동취약 읍면동 분포'); cr.font.size = Pt(8); cr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
rc.paragraphs[0].add_run('이동취약 읍면동(기준: 65세 이상 40% 이상, 밀도 1.27개/㎢ 미만) 요약').bold = True
rc.paragraphs[0].runs[0].font.size = Pt(8.8)
for s_ in [f'- **{V["n"]}개**(분석 대상의 {pc(V["n"] / R["n"])}). 면 {V["by_t"].get("면",0)}개, 읍 {V["by_t"].get("읍",0)}개, 동 {V["by_t"].get("동",0)}개. 면 지역의 {pc(V["mian_vul_ratio"])}가 해당.',
           f'- 거주 인구 {f0(V["pop"])}명(전체의 {pc(V["pop_share"])}), 이 중 65세 이상 {f0(V["old"])}명(전체 고령자의 {pc(V["old_share_of_all_old"])}). 취약 읍면동의 65세 이상 비중 중앙값 {pc(V["median_share"])}.',
           f'- 면적은 {f0(V["area"])}㎢로 분석 대상 면적의 {pc(V["area_share"])}. 정류장 1개당 면적 {V["area_ha_per_stop"]:.0f}ha(동 지역의 {V["area_ha_per_stop"] / T.loc["동","정류장당면적ha"]:.0f}배). 정류장이 0개인 곳 {V["zero"]}개.',
           f'- 시도별: ' + ', '.join(f'{r.sido.replace("특별자치도","").replace("광역시","")} {int(r.vn)}' for r in S.head(7).itertuples()) + f'. 상위 6개 도(경북·경남·전남·전북·충북·강원)가 {pc(S.head(6).vn.sum() / S.vn.sum(), 0)}를 차지하며 서울·부산·대전·광주·제주는 0개.',
           '- 공간적으로 태백·소백산맥 일대 내륙 산간 시군과 남해안 도서 지역에 연속적으로 분포.']:
    pp = rc.add_paragraph(); pp.paragraph_format.space_after = Pt(1.5)
    for i, tx in enumerate(s_.split('**')):
        if tx:
            rr = pp.add_run(tx); rr.bold = (i % 2 == 1); rr.font.size = Pt(8.4); rr.font.name = FONT; rr._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
sr = [['정류장 서비스권 가정', '기준 밀도', '취약 읍면동', '거주 인구', '65세 이상']]
for s_ in sens:
    sr.append([f'반경 {int(s_["r"] * 1000)}m' + (' (본 분석)' if s_['r'] == .5 else ''), f'{s_["lim"]:.2f}', f0(s_['n']), f0(s_['pop']), f0(s_['old'])])
para('', after=1)
table(sr, [3.6, 2.2, 2.6, 2.8, 2.8], size=7.8)
para(f'[표] 임계값 민감도: 서비스권 반경을 300m로 좁게 가정하면 취약 읍면동은 {f0(sens[0]["n"])}개(65세 이상 {f0(sens[0]["old"])}명)로 늘고, 500m에서도 {V["n"]}개가 남음. 어느 기준에서도 취약 읍면동의 대부분은 면 지역이며 분포 양상은 동일.', size=8.2, before=1, after=2)

h2('다. 시군 단위 집중 - 운영 단위 도출')
fig(f'{OUT}/fig4_sigungu.png', 11.4, '[그림 4] 이동취약 읍면동 수 상위 15개 시·군', after=1)
g5 = G[G.vn >= 5]
para(f'- 취약 읍면동이 있는 시군은 {R["sg_count"]}곳이며, 이 중 **{R["sg_ge5"]}개 시군(5개 이상 보유)이 전체 취약 읍면동의 {pc(R["sg_ge5_share"], 0)}**를 포함. 합천군(17개 읍면 중 16개), 의성군(18개 중 14개), 거창군(12개 중 11개)은 관내 대부분 읍면이 해당.', size=8.8)
para('- 읍면 단위로 개별 대응하기보다 시군 단위로 묶어 운영하는 것이 효율적인 지역이 다수임을 보여 줌(자세한 시군별 표는 code/ 산출물 table_sigungu.csv).', size=8.8)

h1('7. 정책 제안 및 기대효과')
h2('가. 제안')
b1(f'**DRT 우선 도입 권역 지정**: 이동취약 읍면동이 5개 이상인 {R["sg_ge5"]}개 시군을 1순위 권역으로 선정(취약 읍면동의 {pc(R["sg_ge5_share"], 0)} 포괄). 합천·의성·거창·산청·진안·곡성·임실처럼 관내 취약 비율이 높은 군은 **군 전체를 하나의 DRT 운영 구역**으로 설계.')
b1(f'**정류장 신설보다 기존 정류장 활용**: 취약 읍면동에도 정류장 {f0(V["stops"])}개가 이미 존재. 이를 호출·승하차 지점과 읍면 소재지 환승 거점으로 활용하면 면적 대비 정류장이 적은 구조적 한계를 시설 신설 없이 보완 가능.')
b1('**정류장 데이터 고도화**: 현행 자료는 연 1회 수집되며 운행빈도·시간대 정보가 없고 일부 시군은 누락됨(본 분석에서 15개 시군구 제외). 정류장 위치와 함께 운행 횟수를 표준 필드로 공개하고 갱신 주기를 단축하면 "정류장은 있으나 버스가 오지 않는" 지역까지 식별 가능.')
b1('**우선순위 지표 표준화**: 인구 대비 정류장 수만으로는 농촌 격차가 가려지므로 "고령비중 + 면적당 정류장 밀도"를 함께 쓰는 지표를 교통약자 이동편의 평가에 병행 적용.')
h2('나. 기대효과 및 적용 사례')
b1(f'**1차 수혜 대상이 명확함**: 취약 읍면동 {V["n"]}개의 65세 이상 {f0(V["old"])}명(전체 고령자의 {pc(V["old_share_of_all_old"])})이 우선 대상. 면적의 {pc(V["area_share"])}를 차지하는 지역에 인구 {pc(V["pop_share"])}가 분산되어 있어 고정 노선 증설보다 호출형이 적합.')
b1(f'**구체적 적용 사례**: 합천군은 취약 읍면 16곳에 65세 이상 {f0(G.iloc[0].vold)}명, 의성군은 14곳에 {f0(G.iloc[1].vold)}명이 거주. 군 단위 DRT 구역 1개로 각각 이들을 포괄할 수 있음.')
b1('**탄소중립 연계(정성적 시사점)**: 고정 노선 증차 대비 수요가 있을 때만 운행하는 호출형은 빈 차 운행을 줄일 수 있고, 승용차 외 대안이 없는 지역에 대중교통 수단을 제공함. 다만 본 분석은 이용 실적·배출량 자료가 없어 **감축량은 산정하지 않았으며**, 도입 후 이용 데이터로 검증이 필요함.')
b1('**활용 방안**: 결과 파일(emd_final.csv, table_sigungu.csv)은 지자체·국토교통부가 DRT 대상지 선정, 교통약자 이동편의 증진계획 수립, 정류장 데이터 보완 우선순위 결정에 바로 쓸 수 있음.')

h1('8. 한계 및 추가 사항 (별도 제목)')
b1('정류장 수는 서비스 수준(운행 빈도)이 아닌 공급 시설만 반영하며, 읍면동 평균 밀도라서 마을 단위 거주지와 정류장 간 실제 보행거리는 반영하지 못함. 산지 비율이 높은 면은 밀도가 낮게 나오는 경향이 있어 해석에 유의가 필요(민감도 분석으로 보완).')
b1('정류장 자료는 지자체 BIS 수집분으로 일부 시군구가 누락되어 제외했고, 제외된 지역의 취약 여부는 판단하지 못함. 면적은 경위도 경계에 대한 근사 계산이며, 행정동 경계(2025.10.01)와 인구(2025.10), 정류장(2025.10.31) 간 시점 차이는 1개월 이내임. 이름 불일치로 결합에서 빠진 읍면동의 인구는 원자료 총인구의 약 1.3%.')

h1('9. 분석도구 및 참고문헌')
b1('분석도구: Python 3 (pandas, shapely, scipy, matplotlib, python-docx), 보고서 변환 LibreOffice. 코드 공개: ' + URL + ' (00_download.py → 01_prepare.py → 02_analysis.py → 03_build_report.py 순서로 실행, README 참조).')
b1('국토교통부, 「전국 버스정류장 위치정보」, 공공데이터포털, https://www.data.go.kr/data/15067528/fileData.do (정보수집일 2025.10.31, 2025.12.09 갱신본).')
b1('행정안전부, 「주민등록 인구통계 - 연령별 인구현황(읍면동, 2025년 10월)」, https://jumin.mois.go.kr/ageStatMonth.do.')
b1('vuski, 「admdongkor - 대한민국 행정동 경계(ver20251001)」, https://github.com/vuski/admdongkor (원자료: 통계청 통계지리정보서비스 SGIS, 가공물 CC BY 4.0).')
b1('재단법인 숲과나눔·한겨레, 「AI와 함께하는 교통문제 해결을 위한 데이터 분석 공모전」 공고문(2026.08.18), https://koreashe.org/board/?mode=view&board_id=22&post_id=95427 (수송부문 온실가스 중 도로교통 비중 약 97% 인용).')
doc.save(DST)
print('saved', DST)
